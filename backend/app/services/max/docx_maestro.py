"""Family-safe DOCX → maestro + resumen.

Juan (Max-e) asks chat to read an attached Word file, emit a master
(maestro) and a summary (resumen), check pagination/TOC, and publish
links. MiniMax used to call ``execute_code`` + LibreOffice and then
dead-end. This path is pure python-docx. LibreOffice is optional: if
it is missing we still deliver both documents and say plainly that
exact pagination was not verified.

No paid APIs. No arbitrary code execution.
"""
from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger("max.docx_maestro")

TOOL_NAME = "process_docx"
MISSING_TOOL_ES = (
    "Esa herramienta no está disponible aquí. "
    "Para un Word adjunto usa process_docx."
)
PAGINATION_UNVERIFIED_ES = (
    "No pude verificar la paginación exacta ni la tabla de contenido "
    "con LibreOffice (no está en este runtime)."
)
NO_FILE_ES = (
    "No encontré el DOCX adjunto. Súbelo otra vez al chat e inténtalo de nuevo."
)
FAILED_ES = "No pude leer ese documento Word. Inténtalo de nuevo en un momento."

# Need a Word file in play (attachment or ".docx" in the text) plus a
# maestro/resumen/pagination ask. Bare "resumen", "procesar", "publica"
# or "enlaces" must not steal agenda / send-me turns.
_DOCX_PROCESS = re.compile(
    r"(?:"
    r"(?:maestro.{0,80}resumen|resumen.{0,80}maestro)|"
    r"(?:maestro|resumen).{0,40}(?:\.docx|\bword\b|documento)|"
    r"(?:\.docx|\bword\b|documento).{0,40}(?:maestro|resumen)|"
    r"\bpaginaci[oó]n\b.{0,40}(?:docx|word|maestro|documento)|"
    r"(?:docx|word|maestro|documento).{0,40}\bpaginaci[oó]n\b|"
    r"\btabla\s+de\s+contenido\b|"
    r"\btable\s+of\s+contents\b|"
    r"proces(?:a|ar)\s+(?:el\s+|este\s+|the\s+)?(?:docx|documento(?:\s+word)?|word|archivo\s+\.docx)|"
    r"lee(?:r)?\s+(?:el\s+|este\s+)?(?:docx|documento\s+word|archivo\s+\.docx)|"
    r"process\s+(?:the\s+)?(?:attached\s+)?(?:docx|word\s+document)"
    r")",
    re.I,
)
_DOCX_NAME = re.compile(r"[\w.\- ()]+\.docx?\b", re.I)
_DOCX_TOKEN = re.compile(r"[\w.\-()]+\.docx?\b", re.I)
_SENTENCE = re.compile(r"(?s).+?(?:[.!?…](?:\s+|$)|$)")


def attachment_is_docx(filename: Optional[str]) -> bool:
    name = (filename or "").split("?")[0].strip()
    return Path(name).suffix.lower() in {".docx", ".doc"}


def is_docx_process_ask(message: Optional[str], filename: Optional[str] = None) -> bool:
    """True when Juan wants maestro/resumen from a Word file — not a generic ask."""
    text = message or ""
    if not text.strip() or not _DOCX_PROCESS.search(text):
        return False
    return attachment_is_docx(filename) or bool(_DOCX_NAME.search(text))


def filename_from_code_job(tool_call: Optional[dict]) -> str:
    """Best-effort .docx name from a leaked execute_code call."""
    if not isinstance(tool_call, dict):
        return ""
    for key in ("filename", "file", "path", "source_path"):
        raw = str(tool_call.get(key) or "").strip()
        if raw:
            return Path(raw).name
    blob = " ".join(str(v) for v in tool_call.values())
    tokens = _DOCX_TOKEN.findall(blob)
    return tokens[-1] if tokens else ""


def looks_like_docx_code_job(tool_call: Optional[dict]) -> bool:
    """True when a hallucinated execute_code call is really a DOCX job."""
    if not isinstance(tool_call, dict):
        return False
    blob = " ".join(str(v) for v in tool_call.values()).lower()
    keys = ("docx", "maestro", "resumen", "libreoffice", "soffice", "python-docx")
    return any(key in blob for key in keys)


def unknown_tool_reply(tool_name: Optional[str] = None) -> str:
    """Short Spanish line for a tool that is not registered here."""
    name = (tool_name or "").strip()
    if name in {"execute_code", "functions.execute_code"}:
        return MISSING_TOOL_ES
    return "Esa herramienta no está disponible aquí. Inténtalo de otra forma."


def _output_dir() -> Path:
    try:
        from app.edition import is_family_edition, require_data_root

        if is_family_edition():
            path = require_data_root() / "uploads" / "documents"
            path.mkdir(parents=True, exist_ok=True)
            return path
    except Exception:
        pass
    from app.services.data_paths import data_root

    path = data_root() / "uploads" / "documents"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _upload_roots() -> list[Path]:
    from app.services.uploads.chat_uploads import upload_search_roots

    roots: list[Path] = []
    for raw in upload_search_roots():
        try:
            roots.append(Path(raw).expanduser().resolve())
        except OSError:
            continue
    return roots


def _under_upload_roots(path: Path, roots: Optional[list[Path]] = None) -> bool:
    try:
        resolved = path.expanduser().resolve()
    except OSError:
        return False
    for root in (roots or _upload_roots()):
        try:
            resolved.relative_to(root)
            return True
        except ValueError:
            continue
    return False


def resolve_source(
    filename: Optional[str],
    explicit_path: Optional[str] = None,
    *,
    path: Optional[str] = None,
) -> Optional[Path]:
    """Resolve a Word file only under this instance's upload_search_roots()."""
    from app.services.uploads.chat_uploads import resolve_uploaded_file

    explicit_path = explicit_path or path
    roots = _upload_roots()
    if not roots:
        return None
    if explicit_path:
        candidate = Path(explicit_path).expanduser()
        try:
            if candidate.is_file() and _under_upload_roots(candidate, roots):
                return candidate.resolve()
        except OSError:
            pass
    name = Path((filename or explicit_path or "").strip()).name
    if not name:
        return None
    found = resolve_uploaded_file(name, roots=roots)
    if found is not None and _under_upload_roots(found, roots):
        return found.resolve()
    return None


def _paragraphs_and_outline(document) -> tuple[list[str], list[str]]:
    parts: list[str] = []
    outline: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = ""
        try:
            style = str(getattr(para.style, "name", "") or "")
        except Exception:
            style = ""
        if style.lower().startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            outline.append(f"{level}. {text}")
            parts.append(text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return parts, outline


def _first_sentences(text: str, limit: int = 2) -> str:
    chunks = []
    for match in _SENTENCE.finditer((text or "").strip()):
        piece = match.group(0).strip()
        if piece:
            chunks.append(piece)
        if len(chunks) >= limit:
            break
    return " ".join(chunks)


def extractive_summary(parts: list[str], outline: list[str], *, max_chars: int = 1800) -> str:
    """Free extractive resumen — headings plus opening sentences. No LLM."""
    blocks: list[str] = []
    if outline:
        blocks.append("Índice (títulos del documento):")
        blocks.extend(f"- {item}" for item in outline[:20])
        blocks.append("")
    intro = next((p for p in parts if p not in {item.split(". ", 1)[-1] for item in outline}), "")
    if intro:
        blocks.append(_first_sentences(intro, 2))
    used = {intro}
    for item in outline[:8]:
        title = item.split(". ", 1)[-1]
        used.add(title)
        after = ""
        try:
            idx = parts.index(title)
            after = next((p for p in parts[idx + 1:] if p not in used), "")
        except ValueError:
            after = ""
        if after:
            used.add(after)
            blocks.append(f"{title}: {_first_sentences(after, 1)}")
        else:
            blocks.append(title)
    if not outline:
        for para in parts[:6]:
            if para not in used:
                blocks.append(_first_sentences(para, 1))
                used.add(para)
    text = "\n".join(block for block in blocks if block is not None).strip()
    if len(text) > max_chars:
        text = text[:max_chars].rsplit("\n", 1)[0] + "\n…"
    return text or "El documento no tenía texto extraíble."


def _write_docx(path: Path, title: str, body_lines: list[str], note: str = "") -> None:
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)
    if note:
        document.add_paragraph(note)
    for line in body_lines:
        text = (line or "").rstrip()
        if not text:
            document.add_paragraph("")
            continue
        if text.startswith("- "):
            document.add_paragraph(text[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", text):
            document.add_paragraph(text, style="List Number")
        else:
            document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def libreoffice_available() -> bool:
    return bool(shutil.which("soffice") or shutil.which("libreoffice"))


def verify_pagination(source: Path) -> dict[str, Any]:
    """Best-effort page/TOC check via LibreOffice. Never required for delivery."""
    result = {
        "available": False,
        "verified": False,
        "pages": None,
        "toc_verified": False,
        "method": "none",
        "note": PAGINATION_UNVERIFIED_ES,
    }
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if not binary:
        return result
    result["available"] = True
    result["method"] = "libreoffice"
    try:
        with tempfile.TemporaryDirectory(prefix="maxe-docx-") as tmp:
            converted = subprocess.run(
                [binary, "--headless", "--nologo", "--nolockcheck", "--convert-to", "pdf", "--outdir", tmp, str(source)],
                capture_output=True,
                text=True,
                timeout=45,
                check=False,
            )
            pdfs = list(Path(tmp).glob("*.pdf"))
            if converted.returncode != 0 or not pdfs:
                result["note"] = (
                    "LibreOffice está instalado pero no pude convertir el DOCX a PDF "
                    "para contar páginas."
                )
                return result
            pages = _pdf_page_count(pdfs[0])
            result["pages"] = pages
            result["verified"] = pages is not None and pages > 0
            result["toc_verified"] = False
            if result["verified"]:
                result["note"] = (
                    f"Paginación aproximada: {pages} página(s) vía LibreOffice. "
                    "La tabla de contenido no se recalculó."
                )
            return result
    except Exception as exc:
        logger.warning("LibreOffice pagination skipped: %s", type(exc).__name__)
        result["note"] = PAGINATION_UNVERIFIED_ES
        return result


def _pdf_page_count(path: Path) -> Optional[int]:
    try:
        data = path.read_bytes()
    except OSError:
        return None
    # Count page objects without a paid/extra PDF stack.
    count = len(re.findall(br"/Type\s*/Page\b", data))
    return count or None


def _safe_stem(name: str) -> str:
    stem = Path(name or "documento").stem
    cleaned = re.sub(r"[^\w\-]+", "_", stem, flags=re.UNICODE).strip("_")
    return (cleaned or "documento")[:60]


def process_docx(
    filename: Optional[str] = None,
    *,
    path: Optional[str] = None,
    title: Optional[str] = None,
) -> dict[str, Any]:
    """Read a Word file and write maestro + resumen into this instance's uploads."""
    source = resolve_source(filename, path=path)
    if source is None:
        return {"ok": False, "error": NO_FILE_ES, "user_reply": NO_FILE_ES}

    try:
        import docx

        document = docx.Document(str(source))
    except Exception as exc:
        logger.warning("process_docx read failed: %s", type(exc).__name__)
        return {"ok": False, "error": FAILED_ES, "user_reply": FAILED_ES}

    parts, outline = _paragraphs_and_outline(document)
    if not parts:
        return {
            "ok": False,
            "error": "El documento está vacío o no tiene texto extraíble.",
            "user_reply": "El documento está vacío o no tiene texto extraíble.",
        }

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = _safe_stem(title or source.name)
    out_dir = _output_dir()
    maestro_name = f"{stem}_maestro_{stamp}.docx"
    resumen_name = f"{stem}_resumen_{stamp}.docx"
    maestro_path = out_dir / maestro_name
    resumen_path = out_dir / resumen_name

    pagination = verify_pagination(source)
    maestro_note = (
        f"Cuerpo completo leído de {source.name}. "
        f"{pagination.get('note') or PAGINATION_UNVERIFIED_ES}"
    )
    outline_lines = ["Índice (títulos detectados):"] + [f"- {item}" for item in outline]
    if not outline:
        outline_lines = [
            "Índice: el DOCX no tenía estilos de título; se lista el cuerpo en orden.",
        ]
    maestro_body = outline_lines + [""] + parts
    summary = extractive_summary(parts, outline)
    resumen_body = [summary]

    _write_docx(maestro_path, f"Maestro — {source.stem}", maestro_body, note=maestro_note)
    _write_docx(resumen_path, f"Resumen — {source.stem}", resumen_body, note=pagination.get("note") or "")

    maestro_url = f"/api/v1/files/view/documents/{maestro_name}"
    resumen_url = f"/api/v1/files/view/documents/{resumen_name}"
    user_reply = format_user_reply(
        source_name=source.name,
        maestro_url=maestro_url,
        resumen_url=resumen_url,
        pagination=pagination,
        paragraph_count=len(parts),
    )
    return {
        "ok": True,
        "source": source.name,
        "paragraph_count": len(parts),
        "outline": outline,
        "pagination": pagination,
        "pagination_verified": bool(pagination.get("verified")),
        "libreoffice": bool(pagination.get("available")),
        "maestro": {
            "filename": maestro_name,
            "path": str(maestro_path),
            "url": maestro_url,
            "bytes": maestro_path.stat().st_size,
        },
        "resumen": {
            "filename": resumen_name,
            "path": str(resumen_path),
            "url": resumen_url,
            "bytes": resumen_path.stat().st_size,
        },
        "user_reply": user_reply,
        "instruction": (
            "The documents are already written. Publish these two links. "
            "Do not say you still need to process the file. "
            "Do not call execute_code. "
            "If pagination was not verified, say that plainly in Spanish."
        ),
    }


def format_user_reply(
    *,
    source_name: str,
    maestro_url: str,
    resumen_url: str,
    pagination: dict[str, Any],
    paragraph_count: int,
) -> str:
    pages = pagination.get("pages")
    if pagination.get("verified") and pages:
        extra = f"Paginación aproximada: {pages} página(s)."
    else:
        extra = PAGINATION_UNVERIFIED_ES
    return (
        f"Listo. Leí el cuerpo completo de {source_name} "
        f"({paragraph_count} bloque(s) de texto) y generé el maestro y el resumen.\n\n"
        f"- Maestro: {maestro_url}\n"
        f"- Resumen: {resumen_url}\n\n"
        f"{extra}"
    )


def prelookup_calls(message: Optional[str], filename: Optional[str] = None) -> list[dict]:
    if not is_docx_process_ask(message, filename):
        return []
    call: dict[str, Any] = {"tool": TOOL_NAME}
    if filename:
        call["filename"] = filename
    named = _DOCX_NAME.search(message or "")
    if named and "filename" not in call:
        call["filename"] = named.group(0)
    return [call]


def prelookup_message(results: list[tuple[str, str]]) -> str:
    if not results:
        return ""
    parts = [
        "Looked up before you answer (real DOCX result from this instance; "
        "use it, do not say you still need to process the file, do not call execute_code):"
    ]
    for tool, body in results:
        parts.append(f"[{tool}] {body}")
    parts.append(
        "Publish the maestro and resumen links. If pagination was not verified, "
        "say that plainly in Spanish. Never print tool-call markup."
    )
    return "\n\n".join(parts)


def looks_like_plan_or_deadend(text: Optional[str]) -> bool:
    """Model promised to do it / refused without LibreOffice / leaked a plan."""
    low = (text or "").strip().lower()
    if not low:
        return False
    markers = (
        "esta vez voy a",
        "voy a ejecutar",
        "no a describir un plan",
        "no voy a publicar",
        "libreoffice no está",
        "no está disponible en este runtime",
        "execute_code",
        "errortype=",
        "tool \"execute_code\"",
        "buscaré el docx",
        "verificaré la paginación",
    )
    return any(marker in low for marker in markers)


def register_tools(tool, ToolResult) -> None:
    @tool(TOOL_NAME)
    def _process_docx(params: dict, desk: Optional[str] = None):
        filename = str(params.get("filename") or params.get("file") or params.get("path") or "").strip()
        explicit = str(params.get("source_path") or "").strip() or None
        title = str(params.get("title") or "").strip() or None
        result = process_docx(filename or None, path=explicit, title=title)
        if not result.get("ok"):
            return ToolResult(tool=TOOL_NAME, success=False, error=result.get("error") or FAILED_ES, result=result)
        return ToolResult(tool=TOOL_NAME, success=True, result=result)


def docx_maestro_doc() -> str:
    return """
### Documentos Word (DOCX)
- **process_docx** — OBLIGATORIO cuando hay un .docx/.doc adjunto y piden maestro,
  resumen, paginación, tabla de contenido o publicar enlaces. Lee el cuerpo completo
  con python-docx, escribe maestro + resumen, y publica los enlaces.
  Si LibreOffice no está, IGUAL entrega los dos documentos y di claramente qué no
  se pudo verificar (paginación exacta / TOC). NUNCA uses execute_code. NUNCA te
  quedes en un plan ni digas "esta vez sí lo hago" sin entregar los enlaces.
  `{"tool": "process_docx", "filename": "archivo.docx"}`
"""


class _DocxMaestroDoc(str):
    def __new__(cls):
        return str.__new__(cls, docx_maestro_doc())

    def __str__(self) -> str:
        return docx_maestro_doc()

    def __add__(self, other):
        return docx_maestro_doc() + other

    def __radd__(self, other):
        return other + docx_maestro_doc()


DOCX_MAESTRO_DOC = _DocxMaestroDoc()
