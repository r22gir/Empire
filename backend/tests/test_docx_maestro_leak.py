"""Max-e ticket fh-bec7a9e9f39e: MiniMax tool-call leak + DOCX fallback.

Juan's Centro de mando chat showed raw MiniMax protocol:

    ]<]minimax[>[<tool_call>
    to=functions.execute_code … Tool "execute_code" is not available

and then dead-ended because LibreOffice was missing.

These tests write only under pytest tmp / cwd. They do not overwrite
tracked files in reports/. No live model. No paid APIs. No real sends.
"""
from __future__ import annotations

import json
from pathlib import Path

from app.routers.max.router import _sanitize_internal_leakage_text
from app.services.max.docx_maestro import (
    extractive_summary,
    format_user_reply,
    is_docx_process_ask,
    looks_like_docx_code_job,
    looks_like_plan_or_deadend,
    process_docx,
    resolve_source,
    unknown_tool_reply,
    verify_pagination,
)
from app.services.max.thinking import strip_protocol_delimiters, strip_thinking
from app.services.max.tool_executor import (
    filter_unregistered_tool_calls,
    parse_tool_blocks,
    strip_tool_blocks,
)

TICKET_LEAK = (
    "Esta vez voy a ejecutar el procesamiento real y no a describir un plan: "
    "buscaré el DOCX adjunto, revisaré su cuerpo completo, generaré el maestro "
    "y el resumen, y verificaré la paginación antes de publicar los enlaces.]"
    "<]minimax[>[<tool_call>\r\n]<]minimax[>[ LibreOffice no está disponible en "
    "este runtime, así que no puedo paginar ni calcular la tabla de contenido "
    "con el método previsto. El DOCX sí queda localizado, pero no voy a "
    "publicar un maestro sin esas comprobaciones.]<]minimax[>[<tool_call>\r\n"
    'to=functions.execute_code code errorType=unknown_error, error=Tool '
    '\\"execute_code\\" is not available in this environment.'
)


def _sample_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Informe de prueba AMP", level=1)
    document.add_paragraph(
        "Este es el cuerpo completo del documento de Juan. "
        "Describe el plan de trabajo de la semana y los acuerdos del equipo."
    )
    document.add_heading("Objetivos", level=2)
    document.add_paragraph(
        "Cerrar el maestro y el resumen sin depender de LibreOffice. "
        "Publicar enlaces reales para el centro de mando."
    )
    document.add_heading("Notas", level=2)
    document.add_paragraph("La paginación exacta es opcional cuando no hay soffice.")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def test_ticket_leak_tokens_are_stripped():
    cleaned = _sanitize_internal_leakage_text(TICKET_LEAK)
    assert "]<]minimax[>" not in cleaned
    assert "<tool_call>" not in cleaned
    assert "to=functions" not in cleaned
    assert "execute_code" not in cleaned
    assert "errorType" not in cleaned
    assert "unknown_error" not in cleaned
    assert "is not available in this environment" not in cleaned
    assert "maestro" in cleaned.lower()
    assert "LibreOffice no está disponible" in cleaned


def test_line_strips_do_not_eat_real_spanish():
    real = (
        "Juan dijo que no use to=functions.execute_code en el chat. "
        "errorType=unknown_error no es un mensaje para él."
    )
    cleaned = _sanitize_internal_leakage_text(real)
    assert "Juan dijo" in cleaned
    assert "no es un mensaje" in cleaned
    visible = strip_tool_blocks(real)
    assert "Juan dijo" in visible


def test_strip_thinking_keeps_spanish_and_drops_delimiters():
    cleaned = strip_thinking(TICKET_LEAK)
    assert "maestro" in cleaned.lower()
    assert "]<]minimax[>" not in cleaned
    assert strip_protocol_delimiters("hola]<]minimax[>[mundo") == "holamundo"


def test_minimax_execute_code_markup_is_stripped_not_looped():
    calls = parse_tool_blocks(TICKET_LEAK)
    assert calls == []
    visible = strip_tool_blocks(TICKET_LEAK)
    assert "<tool_call>" not in visible
    assert "to=functions" not in visible
    assert "execute_code" not in visible
    assert "errorType" not in visible
    leftover = filter_unregistered_tool_calls(
        [{"tool": "execute_code"}],
        [{"tool": "process_docx", "success": True}],
    )
    assert leftover == []


def test_family_missing_tool_is_spanish(monkeypatch, tmp_path):
    monkeypatch.setenv("EMPIRE_EDITION", "amp")
    monkeypatch.setenv("EMPIRE_DATA_DIR", str(tmp_path))
    from app.edition import apply_amp_process_paths
    from app.services.max.tool_executor import ToolResult, execute_tool

    apply_amp_process_paths()
    result = execute_tool({"tool": "execute_code", "code": "print(1)"}, founder=True)
    assert isinstance(result, ToolResult)
    assert result.success is False
    assert result.tool == "execute_code"
    assert result.error == unknown_tool_reply("execute_code")
    assert "is not available in this environment" not in (result.error or "")
    assert "Available tools:" not in (result.error or "")


def test_execute_code_docx_job_redirects_and_succeeds(monkeypatch, tmp_path):
    monkeypatch.setenv("EMPIRE_EDITION", "amp")
    monkeypatch.setenv("EMPIRE_DATA_DIR", str(tmp_path))
    from app.edition import apply_amp_process_paths
    from app.services.max.tool_executor import execute_tool

    apply_amp_process_paths()
    src = _sample_docx(tmp_path / "uploads" / "documents" / "plan.docx")
    assert looks_like_docx_code_job({"tool": "execute_code", "code": f"read {src.name} and write maestro"})
    result = execute_tool(
        {"tool": "execute_code", "code": f"read {src.name} and write maestro"},
        founder=True,
    )
    assert result.success is True
    assert result.tool == "process_docx"
    assert (result.result or {}).get("ok") is True
    assert (result.result or {}).get("_redirected_from") == "execute_code"
    assert Path(result.result["maestro"]["path"]).is_file()
    assert Path(result.result["resumen"]["path"]).is_file()


def test_resolve_source_stays_in_upload_roots(monkeypatch, tmp_path):
    monkeypatch.setenv("EMPIRE_EDITION", "amp")
    monkeypatch.setenv("EMPIRE_DATA_DIR", str(tmp_path))
    from app.edition import apply_amp_process_paths

    apply_amp_process_paths()
    inside = _sample_docx(tmp_path / "uploads" / "documents" / "ok.docx")
    outside = tmp_path / "outside.docx"
    outside.write_bytes(inside.read_bytes())
    assert resolve_source("ok.docx") == inside.resolve()
    assert resolve_source(None, path=str(inside)) == inside.resolve()
    assert resolve_source(None, path=str(outside)) is None
    assert resolve_source(None, path="/etc/passwd") is None


def test_docx_fallback_without_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setenv("EMPIRE_EDITION", "amp")
    monkeypatch.setenv("EMPIRE_DATA_DIR", str(tmp_path))
    from app.edition import apply_amp_process_paths

    apply_amp_process_paths()
    src = _sample_docx(tmp_path / "uploads" / "documents" / "informe_juan.docx")
    monkeypatch.setattr("app.services.max.docx_maestro.libreoffice_available", lambda: False)
    monkeypatch.setattr(
        "app.services.max.docx_maestro.shutil.which",
        lambda *_a, **_k: None,
    )
    pages = verify_pagination(src)
    assert pages["available"] is False
    assert pages["verified"] is False

    result = process_docx(src.name, path=str(src))
    assert result["ok"] is True
    assert result["pagination_verified"] is False
    assert result["libreoffice"] is False
    assert Path(result["maestro"]["path"]).is_file()
    assert Path(result["resumen"]["path"]).is_file()
    assert "/api/v1/files/view/documents/" in result["maestro"]["url"]
    assert "LibreOffice" in result["user_reply"]
    import docx

    master = docx.Document(result["maestro"]["path"])
    body = "\n".join(p.text for p in master.paragraphs)
    assert "Informe de prueba AMP" in body
    assert "Publicar enlaces reales" in body


def test_intent_and_plan_detectors():
    assert is_docx_process_ask(
        "Lee el DOCX, genera el maestro y el resumen, verifica paginación y publica enlaces",
        "informe.docx",
    )
    assert is_docx_process_ask("Genera el maestro y el resumen del DOCX", "adjunto.docx")
    assert not is_docx_process_ask("hola manito", None)
    assert not is_docx_process_ask("resúmeme el día", "informe.docx")
    assert not is_docx_process_ask("publica los enlaces", "informe.docx")
    assert not is_docx_process_ask("procesa esto", None)
    assert looks_like_plan_or_deadend(TICKET_LEAK)
    assert looks_like_plan_or_deadend("") is False
    assert not looks_like_plan_or_deadend("Listo. Aquí están los enlaces del maestro.")


def test_extractive_summary_uses_headings():
    parts = ["Introducción larga del documento.", "Objetivos", "Cerrar el maestro."]
    outline = ["1. Objetivos"]
    summary = extractive_summary(parts, outline)
    assert "Objetivos" in summary


def _family_env(monkeypatch, tmp_path):
    monkeypatch.setenv("EMPIRE_EDITION", "amp")
    monkeypatch.setenv("EMPIRE_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("EMPIRE_DEFAULT_LOCALE", "es")
    monkeypatch.setenv("ASSISTANT_NAME", "Max-e")
    monkeypatch.setenv("EMPIRE_TASK_DB", str(tmp_path / "empire.db"))
    monkeypatch.delenv("FOUNDER_PIN", raising=False)
    from app.edition import apply_amp_process_paths

    apply_amp_process_paths()


def _patch_chat_path(monkeypatch, fake_chat, fake_stream, fake_execute):
    import importlib

    from app.services.max.ai_router import ai_router

    max_router = importlib.import_module("app.routers.max.router")
    dummy = type("Handoff", (), {
        "is_drawing_intent": False,
        "ready": False,
        "missing": [],
        "intent_mode": "unknown",
    })()

    async def no_link(*_a, **_k):
        return None

    monkeypatch.setattr(max_router, "_is_family_edition", lambda: True)
    monkeypatch.setattr(max_router, "_family_pre_search_open", lambda _desk: False)
    monkeypatch.setattr(max_router, "_link_intelligence_response", no_link)
    monkeypatch.setattr(max_router, "_maybe_handle_direct_route_request", lambda *_a, **_k: None)
    monkeypatch.setattr(max_router, "_maybe_handle_gpu_safety_request", lambda *_a, **_k: None)
    monkeypatch.setattr(max_router, "should_run_whats_new_summary", lambda *_a, **_k: False)
    monkeypatch.setattr(max_router, "should_force_runtime_truth_check", lambda *_a, **_k: False)
    monkeypatch.setattr(max_router, "should_clarify_inventory_request", lambda *_a, **_k: False)
    monkeypatch.setattr(max_router, "build_drawing_handoff", lambda *_a, **_k: dummy)
    monkeypatch.setattr(max_router, "execute_tool", fake_execute)
    monkeypatch.setattr(ai_router, "chat", fake_chat)
    monkeypatch.setattr(ai_router, "chat_stream", fake_stream)
    try:
        from app.services.max import factual_guard
        monkeypatch.setattr(factual_guard, "is_factual_question", lambda *_a, **_k: False)
    except Exception:
        pass
    try:
        import app.services.max.continuity_compaction as continuity
        monkeypatch.setattr(continuity, "should_handle_continuity_command", lambda *_a, **_k: False)
        monkeypatch.setattr(continuity, "should_run_continuity_audit", lambda *_a, **_k: False)
    except Exception:
        pass
    return max_router


def test_chat_docx_path_strips_leak_and_delivers(monkeypatch, tmp_path):
    """POST /api/v1/max/chat — the door Juan's Command Center hits.

    Writes proof under tmp_path only. Does not touch reports/.
    """
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.routers.max.router import router
    from app.services.max.ai_router import AIResponse
    from app.services.max.tool_executor import ToolResult

    _family_env(monkeypatch, tmp_path)
    src = _sample_docx(tmp_path / "uploads" / "documents" / "adjunto_juan.docx")
    executed: list[str] = []
    real = process_docx(src.name, path=str(src))
    assert real["ok"] is True

    async def fake_chat(*_a, **_k):
        return AIResponse(content=TICKET_LEAK, model_used="MiniMax-M3")

    async def fake_stream(*_a, **_k):
        yield TICKET_LEAK, "MiniMax-M3"
        return
        yield  # pragma: no cover

    def fake_execute(tool_call, *_a, **_k):
        tool = str((tool_call or {}).get("tool") or "")
        executed.append(tool)
        if tool == "process_docx":
            return ToolResult(tool="process_docx", success=True, result=real)
        if tool == "execute_code":
            return ToolResult(tool="execute_code", success=False, error=unknown_tool_reply(tool))
        return ToolResult(tool=tool or "unknown", success=False, error="unused")

    _patch_chat_path(monkeypatch, fake_chat, fake_stream, fake_execute)
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    client = TestClient(app)
    message = (
        "Procesa el DOCX adjunto: lee el cuerpo completo, genera el maestro "
        "y el resumen, verifica la paginación y publica los enlaces."
    )
    res = client.post(
        "/api/v1/max/chat",
        json={
            "message": message,
            "channel": "dashboard",
            "usage_source": "web",
            "image_filename": src.name,
        },
    )
    assert res.status_code == 200, res.text
    body = res.json()
    text = body["response"]
    proof = tmp_path / "chat_path.json"
    proof.write_text(json.dumps({
        "kind": "offline-test-output",
        "endpoint": "/api/v1/max/chat",
        "ticket": "fh-bec7a9e9f39e",
        "user": message,
        "attachment": src.name,
        "executed_tools": executed,
        "response": text,
        "quality": body.get("quality"),
    }, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    assert executed == ["process_docx"]
    assert "execute_code" not in executed
    assert "]<]minimax[>" not in text
    assert "<tool_call>" not in text
    assert "to=functions" not in text
    assert "execute_code" not in text
    assert "/files/view/documents/" in text
    assert "maestro" in text.lower()
    quality = body.get("quality") or {}
    assert quality.get("label") != "Quality issue detected"
    assert "Quality check flagged" not in text


def test_family_graceful_reply_uses_docx_delivery(monkeypatch, tmp_path):
    import importlib

    _family_env(monkeypatch, tmp_path)
    max_router = importlib.import_module("app.routers.max.router")
    from app.routers.max.router import ChatRequest

    monkeypatch.setattr(max_router, "_is_family_edition", lambda: True)
    req = ChatRequest(
        message="Genera el maestro y el resumen del DOCX",
        image_filename="adjunto.docx",
    )
    delivery = format_user_reply(
        source_name="adjunto.docx",
        maestro_url="/api/v1/files/view/documents/a_maestro.docx",
        resumen_url="/api/v1/files/view/documents/a_resumen.docx",
        pagination={"verified": False, "pages": None},
        paragraph_count=4,
    )
    visible = max_router._family_graceful_visible_reply(
        TICKET_LEAK,
        req,
        [{"tool": "process_docx", "success": True, "result": {"user_reply": delivery}}],
    )
    assert visible == delivery
    other = max_router._family_graceful_visible_reply(
        TICKET_LEAK,
        ChatRequest(message="hola manito"),
        [],
    )
    assert other != unknown_tool_reply("execute_code")
    assert "LibreOffice no está disponible" in other or "maestro" in other.lower()
